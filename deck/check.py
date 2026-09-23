"""เช็คว่ามีอะไรล้นขอบหรือทับกันไหม — เครื่องนี้ไม่มีตัวเรนเดอร์ PPTX

    python deck/check.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Inches

OVERLAP_X, OVERLAP_Y = Inches(.25), Inches(.2)     # เกณฑ์ต้องเป็น Inches ไม่ใช่ Emu


def main():
    p = Presentation(Path(__file__).parent / "henjai.pptx")
    W, H = p.slide_width, p.slide_height
    inr = lambda v: round(Emu(v).inches, 2)
    bad = 0
    for i, s in enumerate(p.slides, 1):
        boxes = []
        for sh in s.shapes:
            if sh.width == W and sh.height == H:    # พื้นหลัง
                continue
            r, b = sh.left + sh.width, sh.top + sh.height
            if sh.left < 0 or sh.top < 0 or r > W + Inches(.02) or b > H + Inches(.02):
                t = sh.text_frame.text[:20].replace("\n", " ") if sh.has_text_frame else ""
                print(f"S{i} ล้นขอบ [{t}] ขวา {inr(r)} ล่าง {inr(b)}")
                bad += 1
            if sh.has_text_frame and sh.text_frame.text.strip():
                boxes.append((sh.left, sh.top, r, b, sh.text_frame.text[:18].replace("\n", " ")))
        for a in range(len(boxes)):
            for c in range(a + 1, len(boxes)):
                x1, y1, x2, y2, t1 = boxes[a]
                X1, Y1, X2, Y2, t2 = boxes[c]
                ox, oy = min(x2, X2) - max(x1, X1), min(y2, Y2) - max(y1, Y1)
                if ox > OVERLAP_X and oy > OVERLAP_Y:
                    print(f'S{i} ทับ {inr(ox)}x{inr(oy)}" [{t1}] <-> [{t2}]')
                    bad += 1
    print("✅ ผ่าน ไม่มีทับซ้อน/ล้นขอบ" if not bad else f"❌ เจอ {bad} จุด")
    print(f"{len(p.slides._sldIdLst)} สไลด์ · {inr(W)} x {inr(H)} นิ้ว")
    return bad


if __name__ == "__main__":
    raise SystemExit(1 if main() else 0)
