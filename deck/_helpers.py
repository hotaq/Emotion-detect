"""สร้างเดค PPTX — แก้เนื้อหาในไฟล์นี้แล้วรันใหม่

    python deck/build.py     ->  deck/henjai.pptx
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).parent
IMG = HERE / "img"

BG = RGBColor(0x0D, 0x0D, 0x0F)
FG = RGBColor(0xFF, 0xFF, 0xFF)
DIM = RGBColor(0x8A, 0x8A, 0x92)
ACC = RGBColor(0xFF, 0xB0, 0x20)          # อำพัน — สื่อความอบอุ่นของคำว่า "ใจ"
TH = "Thonburi"                            # ฟอนต์ไทยที่มีติดเครื่อง mac ทุกเครื่อง

W, H = Inches(13.333), Inches(7.5)


def deck():
    p = Presentation()
    p.slide_width, p.slide_height = W, H
    return p


def slide(p, header=True):
    s = p.slides.add_slide(p.slide_layouts[6])
    bg = s.shapes.add_shape(1, 0, 0, W, H)
    bg.fill.solid(); bg.fill.fore_color.rgb = BG; bg.line.fill.background()
    bg.shadow.inherit = False
    if header:
        txt(s, "TRUE INNOVATION LAUNCHPAD 2026", Inches(.55), Inches(.3),
            Inches(6), Inches(.3), 11, DIM)
        txt(s, "เห็นใจ · HENJAI", Inches(6.8), Inches(.3), Inches(6), Inches(.3),
            11, DIM, align=PP_ALIGN.RIGHT)
    return s


def txt(s, text, x, y, w, h, size, color=FG, bold=False, align=PP_ALIGN.LEFT,
        spacing=1.25, font=TH):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(str(text).split("\n")):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = spacing
        r = para.add_run(); r.text = line
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, font
        r.font.color.rgb = color
    return tb


def pic(s, name, x, y, w=None, h=None):
    f = IMG / name
    if not f.exists():
        box = s.shapes.add_shape(1, x, y, w or Inches(4), h or Inches(3))
        box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0x1C, 0x1C, 0x22)
        box.line.color.rgb = ACC
        box.shadow.inherit = False
        txt(s, f"[ ใส่รูป: {name} ]", x, (y + (h or Inches(3)) // 2),
            w or Inches(4), Inches(.4), 14, ACC, align=PP_ALIGN.CENTER)
        return None
    return s.shapes.add_picture(str(f), x, y, w, h)


def source(s, text):
    txt(s, text, Inches(.55), Inches(6.85), Inches(12.2), Inches(.4), 10, DIM)


def big(s, number, label, x, w, y=2.4, size=72):
    """ตัวเลขใหญ่ + คำอธิบายใต้ — แบบเดียวกับเดค True"""
    txt(s, number, x, Inches(y), w, Inches(1.5), size, ACC, bold=True, align=PP_ALIGN.CENTER)
    txt(s, label, x, Inches(y + 1.55), w, Inches(1.2), 16, FG, align=PP_ALIGN.CENTER)


def todo(s, title, what):
    txt(s, "⚠ ยังไม่มีข้อมูล", Inches(.55), Inches(1.2), Inches(12.2), Inches(.5), 18, ACC, bold=True)
    txt(s, title, Inches(.55), Inches(2.0), Inches(12.2), Inches(.8), 40, FG, bold=True)
    txt(s, what, Inches(.55), Inches(3.2), Inches(12.2), Inches(3), 18, DIM)


