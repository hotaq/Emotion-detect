"""สร้างเดค PPTX — แก้เนื้อหาในไฟล์นี้แล้วรันใหม่

    python deck/build.py     ->  deck/henjai.pptx
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).parent
IMG = HERE / "img"

BG = RGBColor(0x0D, 0x0D, 0x0F)
FG = RGBColor(0xFF, 0xFF, 0xFF)
DIM = RGBColor(0x8A, 0x8A, 0x92)
ACC = RGBColor(0xFF, 0xB0, 0x20)          # อำพัน — สื่อความอบอุ่นของคำว่า "ใจ"
TH = "Thonburi"                            # ฟอนต์ไทยที่มีติดเครื่อง mac ทุกเครื่อง

W, H = Inches(13.333), Inches(7.5)


def deck():
    p = Presentation()
    p.slide_width, p.slide_height = W, H
    return p


def slide(p, header=True):
    s = p.slides.add_slide(p.slide_layouts[6])
    bg = s.shapes.add_shape(1, 0, 0, W, H)
    bg.fill.solid(); bg.fill.fore_color.rgb = BG; bg.line.fill.background()
    bg.shadow.inherit = False
    if header:
        txt(s, "TRUE INNOVATION LAUNCHPAD 2026", Inches(.55), Inches(.3),
            Inches(6), Inches(.3), 11, DIM)
        txt(s, "เห็นใจ · HENJAI", Inches(6.8), Inches(.3), Inches(6), Inches(.3),
            11, DIM, align=PP_ALIGN.RIGHT)
    return s


def txt(s, text, x, y, w, h, size, color=FG, bold=False, align=PP_ALIGN.LEFT,
        spacing=1.25, font=TH):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(str(text).split("\n")):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = spacing
        r = para.add_run(); r.text = line
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, font
        r.font.color.rgb = color
    return tb


def pic(s, name, x, y, w=None, h=None):
    f = IMG / name
    if not f.exists():
        box = s.shapes.add_shape(1, x, y, w or Inches(4), h or Inches(3))
        box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0x1C, 0x1C, 0x22)
        box.line.color.rgb = ACC
        box.shadow.inherit = False
        txt(s, f"[ ใส่รูป: {name} ]", x, (y + (h or Inches(3)) // 2),
            w or Inches(4), Inches(.4), 14, ACC, align=PP_ALIGN.CENTER)
        return None
    return s.shapes.add_picture(str(f), x, y, w, h)


def source(s, text):
    txt(s, text, Inches(.55), Inches(6.85), Inches(12.2), Inches(.4), 10, DIM)


def big(s, number, label, x, w, y=2.4, size=72):
    """ตัวเลขใหญ่ + คำอธิบายใต้ — แบบเดียวกับเดค True"""
    txt(s, number, x, Inches(y), w, Inches(1.5), size, ACC, bold=True, align=PP_ALIGN.CENTER)
    txt(s, label, x, Inches(y + 1.55), w, Inches(1.2), 16, FG, align=PP_ALIGN.CENTER)


def todo(s, title, what):
    txt(s, "⚠ ยังไม่มีข้อมูล", Inches(.55), Inches(1.2), Inches(12.2), Inches(.5), 18, ACC, bold=True)
    txt(s, title, Inches(.55), Inches(2.0), Inches(12.2), Inches(.8), 40, FG, bold=True)
    txt(s, what, Inches(.55), Inches(3.2), Inches(12.2), Inches(3), 18, DIM)


def quote(s, text, size=40, y=2.6, color=FG):
    """ประโยคเดียวกลางจอ — ความสูงคิดจากจำนวนบรรทัดจริง ไม่งั้นกล่องทับกัน"""
    h = .25 + len(text.split("\n")) * size * 1.5 / 72
    txt(s, text, Inches(.8), Inches(y), Inches(11.7), Inches(h), size, color,
        bold=True, align=PP_ALIGN.CENTER, spacing=1.35)


def act(s, name):
    txt(s, name, Inches(.55), Inches(.72), Inches(6), Inches(.3), 10, ACC)


def bullets(s, items, y=2.3, gap=1.15, x=.55, w=12.2):
    for i, (t, d) in enumerate(items):
        yy = Inches(y + i * gap)
        txt(s, t, Inches(x), yy, Inches(w), Inches(.5), 24, ACC, bold=True)
        txt(s, d, Inches(x), yy + Inches(.52), Inches(w), Inches(.6), 17, FG)


# ============================================================
# องก์ 1 · HOOK — 45 วิ
# ============================================================
p = deck()

# 1 ภาพคนคุยกัน ไม่มีตัวหนังสือ
s = slide(p, header=False)
pic(s, "scene_talk.jpg", 0, 0, W, H)
act(s, "องก์ 1 · HOOK")

# 2 คำถาม
s = slide(p, header=False)
quote(s, "เขาเงียบเพราะกำลังคิด", 44, y=2.3, color=DIM)
quote(s, "หรือเพราะกำลังโกรธคุณอยู่", 52, y=3.6)

# 3 เราตอบได้ เขาไม่ได้
s = slide(p, header=False)
quote(s, "เราตอบได้ใน 0.2 วินาที\nเพราะเราเห็นหน้าเขา", 42, y=2.3, color=DIM)
quote(s, "แต่เขาไม่มีทางรู้เลย", 44, y=4.6, color=ACC)

# 4 ตัวเลข
s = slide(p)
txt(s, "ในประเทศไทย", Inches(.55), Inches(1.5), Inches(12.2), Inches(.6), 24, DIM)
txt(s, "169,364", Inches(.55), Inches(2.1), Inches(12.2), Inches(2), 130, ACC, bold=True)
txt(s, "คนพิการทางการเห็น ที่ไม่มีทางรู้ว่าคนตรงหน้ากำลังรู้สึกยังไง",
    Inches(.55), Inches(4.5), Inches(12.2), Inches(.7), 26, FG)
source(s, "ที่มา : กรมส่งเสริมและพัฒนาคุณภาพชีวิตคนพิการ (พก.) ข้อมูล ณ 31 กรกฎาคม 2569")

# 5 AI วันนี้บอกได้แค่นี้
s = slide(p, header=False)
pic(s, "angry.jpg", 0, 0, h=H)
box = s.shapes.add_shape(1, Inches(7.4), Inches(3.0), Inches(5.2), Inches(1.4))
box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0x1C, 0x1C, 0x22)
box.line.color.rgb = DIM; box.shadow.inherit = False
txt(s, "person — 1", Inches(7.4), Inches(3.4), Inches(5.2), Inches(.7), 34, DIM,
    bold=True, align=PP_ALIGN.CENTER)
txt(s, "เทคโนโลยีทุกวันนี้บอกได้แค่นี้", Inches(7.4), Inches(4.7), Inches(5.2),
    Inches(.6), 18, DIM, align=PP_ALIGN.CENTER)

# 6 เราบอกได้มากกว่า — ภาพเดียวกัน
s = slide(p, header=False)
pic(s, "angry.jpg", 0, 0, h=H)
box = s.shapes.add_shape(1, Inches(7.0), Inches(2.8), Inches(5.8), Inches(1.9))
box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0x1C, 0x1C, 0x22)
box.line.color.rgb = ACC; box.shadow.inherit = False
txt(s, "“ชายตรงหน้าคุณกำลังโกรธ\nคิ้วขมวดแน่น”", Inches(7.0), Inches(3.1),
    Inches(5.8), Inches(1.4), 26, ACC, bold=True, align=PP_ALIGN.CENTER)
txt(s, "สิ่งที่เขาต้องการจริง ๆ", Inches(7.0), Inches(5.0), Inches(5.8),
    Inches(.6), 18, FG, align=PP_ALIGN.CENTER)

# 7 ชื่อแบรนด์
s = slide(p, header=False)
txt(s, "เห็นใจ", 0, Inches(2.3), W, Inches(2), 110, FG, bold=True, align=PP_ALIGN.CENTER)
txt(s, "H E N J A I", 0, Inches(4.3), W, Inches(.8), 26, ACC, align=PP_ALIGN.CENTER)
txt(s, "มองไม่เห็น ก็เห็นใจกันได้", 0, Inches(5.3), W, Inches(.8), 24, DIM,
    align=PP_ALIGN.CENTER)

# ============================================================
# องก์ 2 · ปัญหา
# ============================================================
s = slide(p); act(s, "องก์ 2 · ปัญหา")
txt(s, "อ่านสีหน้าคู่สนทนาไม่ได้", Inches(.55), Inches(1.2), Inches(12.2), Inches(.8),
    40, FG, bold=True)
pic(s, "pain_face.jpg", Inches(7.6), Inches(2.2), Inches(5.2))
bullets(s, [("ไม่รู้ว่าเขายิ้ม โกรธ เศร้า หรือเบื่อ", "ได้ยินแค่เสียง ขาดข้อมูลจากใบหน้า"),
            ("ตอบสนองไม่ถูก", "ควรปลอบ ควรหยุดพูด หรือควรถามต่อ"),
            ("เสี่ยงเข้าใจผิด", "คิดว่าเขาไม่พอใจ ทั้งที่แค่ทำหน้าจริงจัง")], y=2.4, w=6.6)

s = slide(p); act(s, "องก์ 2 · ปัญหา")
txt(s, "อ่านภาษากายไม่ได้", Inches(.55), Inches(1.2), Inches(12.2), Inches(.8),
    40, FG, bold=True)
pic(s, "pain_body.jpg", Inches(7.6), Inches(2.2), Inches(5.2))
bullets(s, [("ไม่รู้ว่าเขาหันมาหาเรา", "หรือกำลังคุยกับคนอื่นอยู่"),
            ("ประเมินระยะห่างไม่ได้", "ยืนใกล้เกินไป หรือไกลเกินไป"),
            ("พลาดสัญญาณทางสังคม", "พยักหน้า โบกมือ ชี้นิ้ว กอดอก")], y=2.4, w=6.6)

s = slide(p); act(s, "องก์ 2 · ปัญหา")
txt(s, "ไม่รู้บริบทของสถานที่", Inches(.55), Inches(1.2), Inches(12.2), Inches(.8),
    40, FG, bold=True)
pic(s, "pain_env.jpg", Inches(7.6), Inches(2.2), Inches(5.2))
bullets(s, [("ไม่รู้ว่ามีอะไรขวางอยู่", "โต๊ะ เก้าอี้ ประตู เคาน์เตอร์"),
            ("ไม่รู้ว่าที่นี่ที่ไหน", "ร้านกาแฟ ห้องประชุม หรือสถานีรถไฟ"),
            ("ไม่รู้ลำดับความสำคัญ", "อะไรควรรู้ทันที อะไรเป็นแค่รายละเอียด")], y=2.4, w=6.6)

# ผลที่ตามมา — ตัวเลขจริง
s = slide(p); act(s, "องก์ 2 · ปัญหา")
txt(s, "ผลที่ตามมา", Inches(.55), Inches(1.2), Inches(12.2), Inches(.7), 36, FG, bold=True)
big(s, "17.2%", "ไม่ได้รับการศึกษา\nเลยแม้แต่ระดับเดียว", Inches(.55), Inches(3.9), y=2.5)
big(s, "68.9%", "จบสูงสุด\nแค่ชั้นประถมศึกษา", Inches(4.7), Inches(3.9), y=2.5)
big(s, "1.9%", "จบปริญญาตรี\nหรือเทียบเท่า", Inches(8.85), Inches(3.9), y=2.5)
source(s, "ที่มา : กรมส่งเสริมและพัฒนาคุณภาพชีวิตคนพิการ (พก.) ณ 30 เม.ย. 2569 · เฉพาะผู้พิการทางการเห็น 168,763 คน")

# เคสจริง (ยังไม่มี)
s = slide(p); act(s, "องก์ 2 · ปัญหา")
todo(s, "เคสจริง — ผู้ใช้ 1 คน",
     "ต้องสัมภาษณ์คนตาบอดจริงอย่างน้อย 1 คน ผ่านสมาคมคนตาบอดแห่งประเทศไทย\n\n"
     "ใบนี้คือใบที่แข็งที่สุดของเดคที่ชนะ — ต้องมี ชื่อ อายุ เหตุการณ์จริงที่เขาพลาด\n"
     "เพราะอ่านสีหน้าไม่ออก และพฤติกรรมการใช้มือถือทุกวันนี้")

# เทียบของที่มีอยู่
s = slide(p); act(s, "องก์ 2 · ปัญหา")
txt(s, "ของที่มีอยู่แล้ว บอกได้แค่ไหน", Inches(.55), Inches(1.2), Inches(12.2),
    Inches(.8), 36, FG, bold=True)
txt(s, "ที่มีอยู่", Inches(.55), Inches(2.5), Inches(5.8), Inches(.5), 20, DIM, bold=True)
txt(s, "“มีคนหนึ่งคน\nมีโต๊ะอยู่ทางขวา”", Inches(.55), Inches(3.1), Inches(5.8),
    Inches(2), 30, DIM)
txt(s, "บอกสิ่งที่เห็น", Inches(.55), Inches(5.2), Inches(5.8), Inches(.5), 18, DIM)
txt(s, "เห็นใจ", Inches(7.0), Inches(2.5), Inches(5.8), Inches(.5), 20, ACC, bold=True)
txt(s, "“ชายตรงหน้าคุณกำลังโกรธ\nคิ้วขมวดแน่น”", Inches(7.0), Inches(3.1),
    Inches(5.8), Inches(2), 30, FG, bold=True)
txt(s, "บอกความหมาย", Inches(7.0), Inches(5.2), Inches(5.8), Inches(.5), 18, ACC)

# ============================================================
# องก์ 3 · กลุ่มเป้าหมาย
# ============================================================
s = slide(p); act(s, "องก์ 3 · กลุ่มเป้าหมาย")
txt(s, "เขามองเห็นแค่ไหน", Inches(.55), Inches(1.15), Inches(12.2), Inches(.8),
    38, FG, bold=True)
pic(s, "vision_levels.png", Inches(.55), Inches(2.25), Inches(12.2))
for i, lv in enumerate(["ปกติ", "ระดับ 1-2", "ระดับ 3", "ระดับ 4", "ระดับ 5"]):
    txt(s, lv, Inches(.55 + i * 2.44), Inches(4.25), Inches(2.44), Inches(.45),
        17, ACC if i >= 2 else DIM, bold=(i >= 2), align=PP_ALIGN.CENTER)
txt(s, "เราเจาะกลุ่มระดับ 3 ถึง 5", Inches(.55), Inches(5.0), Inches(12.2), Inches(.7),
    32, FG, bold=True, align=PP_ALIGN.CENTER)
txt(s, "บกพร่องรุนแรงถึงตาบอดสนิท — กลุ่มที่เสียงคือช่องทางเดียวที่เหลืออยู่",
    Inches(.55), Inches(5.75), Inches(12.2), Inches(.6), 19, DIM, align=PP_ALIGN.CENTER)
source(s, "ภาพจำลองเพื่อการสื่อสาร ไม่ใช่ภาพทางการแพทย์ · เกณฑ์ระดับตามมาตรฐานสากลที่อ้างในเอกสารแนวคิด")

# ============================================================
# องก์ 4 · ทางแก้
# ============================================================
s = slide(p); act(s, "องก์ 4 · ทางแก้")
txt(s, "โจทย์", Inches(.55), Inches(1.2), Inches(12.2), Inches(.8), 40, FG, bold=True)
txt(s, "ไม่ใช่แค่บอกว่าเห็นอะไร  แต่ต้องผ่านสี่ข้อนี้ให้ได้",
    Inches(.55), Inches(2.1), Inches(12.2), Inches(.6), 22, DIM)
for i, (t, d) in enumerate([
        ("พูดตอนไหน", "พูดทุกวินาทีคือเสียงรบกวน\nเงียบเกินไปคือพลาดเรื่องสำคัญ"),
        ("ผิดแล้วอันตราย", "บอกว่าเขาอารมณ์ดีทั้งที่กำลังโกรธ\nแย่กว่าไม่บอกอะไรเลย"),
        ("ต้องถือมือถือตลอดไหม", "มือข้างหนึ่งถือไม้เท้า\nอีกข้างต้องว่างไว้"),
        ("เน็ตหลุดแล้วยังใช้ได้ไหม", "รถไฟฟ้า ลิฟต์ ใต้ดิน\nคือที่ที่ต้องการมันที่สุด")]):
    x = Inches(.55 + i * 3.17)
    txt(s, t, x, Inches(3.2), Inches(2.95), Inches(.8), 21, ACC, bold=True)
    txt(s, d, x, Inches(4.1), Inches(2.95), Inches(2), 16, FG)

s = slide(p, header=False); act(s, "องก์ 4 · ทางแก้")
txt(s, "กล้องมือถือที่คุณมีอยู่แล้ว", Inches(.55), Inches(2.0), Inches(12.2),
    Inches(.85), 38, DIM, align=PP_ALIGN.CENTER)
txt(s, "สู่ผู้ช่วยที่บอกได้ว่า\nคนตรงหน้ากำลังรู้สึกยังไง", Inches(.55), Inches(3.0),
    Inches(12.2), Inches(2), 50, FG, bold=True, align=PP_ALIGN.CENTER)
txt(s, "ไม่ต้องซื้ออุปกรณ์เพิ่ม   ·   ไม่ต้องสวมใส่อะไร   ·   พูดเป็นภาษาไทย",
    Inches(.55), Inches(5.5), Inches(12.2), Inches(.6), 20, ACC, align=PP_ALIGN.CENTER)

s = slide(p); act(s, "องก์ 4 · ทางแก้")
txt(s, "สามอย่างที่มันอ่านให้", Inches(.55), Inches(1.2), Inches(12.2), Inches(.8),
    38, FG, bold=True)
for i, (t, d, ex) in enumerate([
        ("อารมณ์", "สีหน้า แววตา คิ้ว", "“กำลังโกรธ คิ้วขมวดแน่น”"),
        ("ภาษากาย", "หันมาทางเรา ระยะห่าง ท่าทาง", "“ยืนกอดอก หันตัวมาทางคุณ”"),
        ("สภาพแวดล้อม", "ที่นี่ที่ไหน มีอะไรขวาง", "“ชานชาลารถไฟ คนแน่นมาก”")]):
    x = Inches(.55 + i * 4.25)
    txt(s, t, x, Inches(2.5), Inches(4), Inches(.7), 30, ACC, bold=True)
    txt(s, d, x, Inches(3.3), Inches(4), Inches(.8), 17, DIM)
    txt(s, ex, x, Inches(4.2), Inches(4), Inches(1.5), 19, FG)
source(s, "ตัวอย่างข้างบนคือผลจริงที่ระบบพูดออกมาระหว่างทดสอบ")

# ============================================================
# องก์ 5 · สถาปัตยกรรม
# ============================================================
def card(s, x, y, w, h, edge=None, fill=0x1C1C22):
    from pptx.enum.shapes import MSO_SHAPE
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor(fill >> 16, (fill >> 8) & 255, fill & 255)
    if edge:
        sh.line.color.rgb = edge; sh.line.width = Pt(1.5)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    sh.adjustments[0] = .12
    return sh


def arrow(s, x, y, w=.45):
    txt(s, "›", x, y, Inches(w), Inches(.6), 40, ACC, bold=True, align=PP_ALIGN.CENTER)


s = slide(p); act(s, "องก์ 5 · สถาปัตยกรรม")
txt(s, "ทำงานยังไง", Inches(.55), Inches(1.1), Inches(12.2), Inches(.8), 38, FG, bold=True)

# กล้อง
card(s, Inches(.55), Inches(2.3), Inches(1.9), Inches(2.5))
pic(s, "ic_cam.png", Inches(1.12), Inches(2.65), Inches(.78), Inches(.78))
txt(s, "กล้อง", Inches(.55), Inches(3.6), Inches(1.9), Inches(.4), 18, FG,
    bold=True, align=PP_ALIGN.CENTER)
txt(s, "30 เฟรม\nต่อวินาที", Inches(.55), Inches(4.02), Inches(1.9), Inches(.7), 13,
    DIM, align=PP_ALIGN.CENTER)
arrow(s, Inches(2.55), Inches(3.25))

# สามสาย
for i, (ic, t, d) in enumerate([("ic_emotion", "อารมณ์", "สีหน้า แววตา คิ้ว"),
                                ("ic_body", "ภาษากาย", "ท่าทาง ระยะห่าง"),
                                ("ic_env", "สภาพแวดล้อม", "สถานที่ สิ่งกีดขวาง")]):
    y = Inches(1.85 + i * 1.25)
    card(s, Inches(3.15), y, Inches(3.5), Inches(1.15), edge=ACC)
    pic(s, ic + ".png", Inches(3.42), y + Inches(.27), Inches(.6), Inches(.6))
    txt(s, t, Inches(4.25), y + Inches(.16), Inches(2.25), Inches(.38), 18, FG, bold=True)
    txt(s, d, Inches(4.25), y + Inches(.58), Inches(2.25), Inches(.36), 12, DIM)
arrow(s, Inches(6.75), Inches(3.25))

# รวม -> จัดลำดับ -> พูด
for i, (ic, t, d) in enumerate([("ic_fuse", "รวมความหมาย", "ต่อเป็นสถานการณ์เดียว"),
                                ("ic_priority", "จัดลำดับ", "อันตรายขึ้นก่อน"),
                                ("ic_speak", "พูดภาษาไทย", "ออกลำโพงทันที")]):
    x = Inches(7.3 + i * 2.05)
    card(s, x, Inches(2.3), Inches(1.7), Inches(2.5))
    pic(s, ic + ".png", x + Inches(.48), Inches(2.62), Inches(.74), Inches(.74))
    txt(s, t, x, Inches(3.55), Inches(1.7), Inches(.4), 15, FG, bold=True, align=PP_ALIGN.CENTER)
    txt(s, d, x, Inches(3.98), Inches(1.7), Inches(.7), 12, DIM, align=PP_ALIGN.CENTER)
    if i < 2:
        arrow(s, x + Inches(1.72), Inches(3.25), w=.33)

txt(s, "“ชายตรงหน้าคุณกำลังโกรธ คิ้วขมวดแน่น”", Inches(.55), Inches(5.7),
    Inches(12.2), Inches(.65), 26, ACC, bold=True, align=PP_ALIGN.CENTER)
txt(s, "ผลจริงที่ระบบพูดออกมาระหว่างทดสอบ", Inches(.55), Inches(6.45), Inches(12.2),
    Inches(.45), 16, DIM, align=PP_ALIGN.CENTER)

s = slide(p); act(s, "องก์ 5 · สถาปัตยกรรม")
txt(s, "ไม่ใช่พูดตลอดเวลา", Inches(.55), Inches(1.2), Inches(12.2), Inches(.8),
    38, FG, bold=True)
txt(s, "เราวัดก่อนว่าภาพเปลี่ยนไปแค่ไหน แล้วค่อยตัดสินใจว่าจะถามหรือไม่ถาม",
    Inches(.55), Inches(2.1), Inches(12.2), Inches(.6), 20, DIM)
for i, (n, t, d) in enumerate([
        ("นิ่ง", "ไม่พูดเลย", "ภาพไม่เปลี่ยน ก็ไม่มีอะไรต้องบอก"),
        ("ขยับนิด", "บอกเฉพาะที่เปลี่ยน", "ไม่ทวนฉากเดิมซ้ำ ๆ"),
        ("เปลี่ยนฉาก", "บรรยายที่ใหม่ทั้งหมด", "เข้าห้องใหม่ เจอคนใหม่"),
        ("กะทันหัน", "เตือนก่อนเลย", "มีอะไรเคลื่อนเข้ามาใกล้")]):
    x = Inches(.55 + i * 3.17)
    txt(s, n, x, Inches(3.0), Inches(2.95), Inches(.6), 26, ACC, bold=True)
    txt(s, t, x, Inches(3.8), Inches(2.95), Inches(.6), 20, FG, bold=True)
    txt(s, d, x, Inches(4.5), Inches(2.95), Inches(1.3), 15, DIM)

# ============================================================
# องก์ 6 · เดโม
# ============================================================
s = slide(p); act(s, "องก์ 6 · เดโม")
txt(s, "หน้าตาของจริง", Inches(.55), Inches(1.2), Inches(12.2), Inches(.8),
    38, FG, bold=True)
pic(s, "app.png", Inches(.55), Inches(2.2), Inches(4.2), Inches(4.2))
txt(s, "ทั้งหน้าจอคือปุ่มเดียว", Inches(5.3), Inches(2.4), Inches(7.5), Inches(.6),
    26, ACC, bold=True)
txt(s, "ผู้ใช้เรามองไม่เห็น จอจึงไม่มีอะไรให้มอง\n"
       "กดตรงไหนก็ได้ แล้วฟังอย่างเดียว\n"
       "เวลามันพูด จะมีคลื่นเสียงขึ้นให้คนข้าง ๆ รู้ว่ากำลังทำงานอยู่",
    Inches(5.3), Inches(3.1), Inches(7.5), Inches(2), 19, FG)

s = slide(p, header=False); act(s, "องก์ 6 · เดโม")
quote(s, "เดโมสด", 60, y=2.8)
txt(s, "หันกล้องไปทางกรรมการ", 0, Inches(4.5), W, Inches(.6), 22, DIM,
    align=PP_ALIGN.CENTER)

# ============================================================
# องก์ 7 · ผลวัด
# ============================================================
s = slide(p); act(s, "องก์ 7 · ผลวัด")
txt(s, "เราวัดจริง ไม่ได้เดโมแต่ทางที่สวย", Inches(.55), Inches(1.2), Inches(12.2),
    Inches(.8), 36, FG, bold=True)
pic(s, "testset.png", Inches(.55), Inches(2.3), Inches(12.2))
txt(s, "ภาพทดสอบ 24 ใบ 4 กลุ่ม — อารมณ์ · ภาษากาย · สภาพแวดล้อม · ภาพที่ไม่มีคน",
    Inches(.55), Inches(4.3), Inches(12.2), Inches(.6), 20, FG)
source(s, "ภาพจาก Openverse เฉพาะสัญญาอนุญาต CC BY / CC BY-SA / CC0 · เก็บที่มาและผู้สร้างครบทุกใบ")

s = slide(p); act(s, "องก์ 7 · ผลวัด")
txt(s, "ผลที่ได้", Inches(.55), Inches(1.2), Inches(12.2), Inches(.7), 36, FG, bold=True)
for i, (n, t) in enumerate([("6 / 6", "อ่านภาษากายถูก"), ("6 / 6", "อ่านสภาพแวดล้อมถูก"),
                            ("4 / 6", "อ่านอารมณ์ถูก"), ("2.2 วิ", "ต่อการบรรยายหนึ่งครั้ง")]):
    x = Inches(.55 + i * 3.17)
    txt(s, n, x, Inches(2.6), Inches(2.95), Inches(1), 44, ACC, bold=True, align=PP_ALIGN.CENTER)
    txt(s, t, x, Inches(3.9), Inches(2.95), Inches(.8), 17, FG, align=PP_ALIGN.CENTER)
txt(s, "อารมณ์คือด้านที่ยังต้องพัฒนาต่อ — และเรารู้แล้วว่าเพราะอะไร",
    Inches(.55), Inches(5.2), Inches(12.2), Inches(.6), 22, DIM, align=PP_ALIGN.CENTER)

# ไม้ตาย
s = slide(p); act(s, "องก์ 7 · ผลวัด")
txt(s, "ความผิดพลาดที่เราเจอเอง", Inches(.55), Inches(1.2), Inches(12.2), Inches(.8),
    36, FG, bold=True)
pic(s, "angry.jpg", Inches(.55), Inches(2.2), h=Inches(4.2))
txt(s, "เวอร์ชันแรกบรรยายว่า", Inches(6.4), Inches(2.4), Inches(6.4), Inches(.5), 18, DIM)
txt(s, "“ชายหนุ่มหน้าตาใจดี\nกำลังส่งยิ้มทะเล้นให้คุณ”", Inches(6.4), Inches(2.9),
    Inches(6.4), Inches(1.3), 26, DIM, bold=True)
txt(s, "หลังแก้ระบบ", Inches(6.4), Inches(4.5), Inches(6.4), Inches(.5), 18, ACC)
txt(s, "“ชายคนหนึ่งกำลังทำหน้าตาโกรธจัด\nคิ้วขมวดแน่นจนหน้าผากย่น”", Inches(6.4),
    Inches(5.0), Inches(6.4), Inches(1.3), 26, ACC, bold=True)
source(s, "ถ้าผู้ใช้ได้ยินประโยคแรก เขาจะเดินเข้าไปคุยต่อ ทั้งที่อีกฝ่ายกำลังโกรธ")

s = slide(p); act(s, "องก์ 7 · ผลวัด")
txt(s, "ทำงานได้จริงบนเครื่องธรรมดา", Inches(.55), Inches(1.2), Inches(12.2),
    Inches(.8), 36, FG, bold=True)
big(s, "30", "เฟรมต่อวินาที\nวิดีโอไม่กระตุก", Inches(.55), Inches(3.9), y=2.6)
big(s, "2.2", "วินาที\nตั้งแต่เห็นจนพูดออกมา", Inches(4.7), Inches(3.9), y=2.6)
big(s, "0", "ครั้งที่ระบบค้าง\nตลอดการทดสอบ", Inches(8.85), Inches(3.9), y=2.6)

# ============================================================
# องก์ 8 · ต่อไป + ปิด
# ============================================================
s = slide(p); act(s, "องก์ 8 · ก้าวต่อไป")
txt(s, "แผนสร้างและทดลอง ภายใน 12 สัปดาห์", Inches(.55), Inches(1.2), Inches(12.2),
    Inches(.8), 36, FG, bold=True)
for i, (wk, t, d) in enumerate([
        ("W1–2", "เก็บข้อมูลผู้ใช้จริง", "สัมภาษณ์ผู้ใช้ 10 คน\nผ่านสมาคมคนตาบอดฯ"),
        ("W3–6", "ต่อโมดูลจริง", "MediaPipe + ตัวจำแนกอารมณ์\nYOLO11-Pose อ่านท่าทาง"),
        ("W7–10", "ทดลองกับผู้ใช้จริง", "ใช้จริง 10 คน 2 สัปดาห์\nวัดความถูกต้องและความพอใจ"),
        ("W11–12", "สรุปและตัดสิน", "ประเมินผล ทบทวนปัญหา\nวางแผนขยาย")]):
    x = Inches(.55 + i * 3.17)
    txt(s, wk, x, Inches(2.6), Inches(2.95), Inches(.5), 20, ACC, bold=True)
    txt(s, t, x, Inches(3.2), Inches(2.95), Inches(.6), 21, FG, bold=True)
    txt(s, d, x, Inches(4.0), Inches(2.95), Inches(2), 16, DIM)
txt(s, "เกณฑ์ที่ถือว่าผ่าน : อ่านอารมณ์ถูก ≥ 80%  ·  บรรยายสิ่งที่ไม่มีจริง ≤ 10%  ·  พูดภายใน ≤ 3 วินาที",
    Inches(.55), Inches(6.1), Inches(12.2), Inches(.6), 18, ACC)

s = slide(p); act(s, "องก์ 8 · ทีม")
txt(s, "ทีม Rushhour", Inches(.55), Inches(1.2), Inches(12.2), Inches(.8), 38, FG, bold=True)
txt(s, "วิศวกรรมคอมพิวเตอร์และเทคโนโลยีดิจิทัล (CEDT) จุฬาลงกรณ์มหาวิทยาลัย",
    Inches(.55), Inches(2.1), Inches(12.2), Inches(.5), 18, DIM)
for i, (name, skills) in enumerate([
        ("กิตติพงศ์ สมสอาด", "— ใส่บทบาท\n— ใส่สกิล\n— ใส่สกิล"),
        ("ชินพัฒน์ ขุนเจริญ", "— ใส่บทบาท\n— ใส่สกิล\n— ใส่สกิล"),
        ("เรข์ณพัศ สุวรรณปริสุทธิ์", "— ใส่บทบาท\n— ใส่สกิล\n— ใส่สกิล")]):
    x = Inches(.55 + i * 4.25)
    txt(s, name, x, Inches(3.3), Inches(4), Inches(.6), 22, FG, bold=True)
    txt(s, skills, x, Inches(4.1), Inches(4), Inches(2), 16, DIM)

# ปิด — ทวนคำถามจากใบที่ 2
s = slide(p, header=False)
quote(s, "เขาเงียบเพราะกำลังคิด\nหรือเพราะกำลังโกรธ", 38, y=1.7, color=DIM)
quote(s, "วันนี้เราตอบให้เขาได้แล้ว", 36, y=4.0)
txt(s, "เห็นใจ — มองไม่เห็น ก็เห็นใจกันได้", 0, Inches(5.4), W, Inches(.8),
    30, ACC, bold=True, align=PP_ALIGN.CENTER)

# appendix
s = slide(p)
txt(s, "สถาปัตยกรรมตามเอกสารแนวคิด", Inches(.55), Inches(1.1), Inches(6), Inches(.7),
    28, FG, bold=True)
pic(s, "architecture.png", Inches(4.6), Inches(1.0), h=Inches(6.0))

s = slide(p)
txt(s, "ที่มาข้อมูล", Inches(.55), Inches(1.2), Inches(12.2), Inches(.7), 32, FG, bold=True)
txt(s, "กรมส่งเสริมและพัฒนาคุณภาพชีวิตคนพิการ (พก.) — รายงานข้อมูลสถานการณ์ด้านคนพิการ\n"
       "    ข้อมูล ณ 31 ก.ค. 2569 · dep.go.th/images/uploads/files/Dashboard072569.pdf\n\n"
       "สำนักงานสถิติแห่งชาติ — การสำรวจความพิการ พ.ศ. 2565 (ร่วมกับ UNICEF)\n"
       "    nso.go.th/nsoweb/storage/survey_detail/2023/20230725162812_14118.pdf\n\n"
       "WHO — Blindness and vision impairment fact sheet (ปรับปรุง 10 ก.พ. 2026)\n"
       "    who.int/news-room/fact-sheets/detail/blindness-and-visual-impairment\n\n"
       "IAPB Vision Atlas — Thailand country data · visionatlas.iapb.org/country-data/thailand/\n\n"
       "ภาพทดสอบทั้งหมดจาก Openverse เฉพาะสัญญาอนุญาตที่ใช้ได้ — รายชื่อผู้สร้างครบใน poc/eval/manifest.json",
    Inches(.55), Inches(2.1), Inches(12.2), Inches(4.5), 15, DIM)

out = HERE / "henjai.pptx"
p.save(out)
print(f"{len(p.slides._sldIdLst)} สไลด์ -> {out}")
